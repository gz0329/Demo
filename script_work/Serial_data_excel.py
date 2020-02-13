from datetime import datetime
from openpyxl import Workbook
from openpyxl.chart import (
    ScatterChart,
    Reference,
    Series,
)

def time_dif( t1, t2 ):
	t_dif = datetime.strptime( t2, '%H:%M:%S') - datetime.strptime( t1, '%H:%M:%S')
	return t_dif.seconds

f = open( "Test_mag_time_temp_humidity.txt", "r")

wb = Workbook()
ws = wb.active

title = ['Time', 'Mag_norm', 'Mag_orth', 'Board_temp', 'Ambient_temp', 'Humidity' ]

Time_buf = []
Mag_norm_buf = []
Mag_orth_buf = []
Board_temp_buf = []
Ambient_temp_buf = []
Humidity_buf = []

line_count = 0
first_line = f.readline()
ref_time = first_line.split()[ 0 ]

f.seek( 0 )

for lines in f:
    if ( line_count < 1000000 ):
        strings = lines.split()

        time_delta = time_dif( ref_time, strings[ 0 ] )
        #print( time_delta )
        Mag_norm_f = float( strings[ 1 ][ 1: -2 ] )
        Mag_orth_f = float( strings[ 2 ][ 1: -2 ] )

        if ( Mag_norm_f < 10000 and Mag_orth_f < 10000 and Mag_orth_f > 1 and Mag_norm_f > 1 ):      	
            Time_buf.append( time_delta )
            Mag_norm_buf.append( strings[ 1 ][ 1: -2 ] )
            Mag_orth_buf.append( strings[ 2 ][ 1: -2 ] )
            Board_temp_buf.append( strings[ 3 ][ 1: -2 ] )
            Ambient_temp_buf.append( strings[ 4 ][ 1: -2 ] )
            Humidity_buf.append( strings[ 5 ][ 1: -2 ] )
            line_count = line_count + 1
        #print( strings[ 0 ] + " " + strings[ 1 ][ 1: -2 ] + " " + strings[ 2 ][ 1: -2 ] )
    else:
        pass

col_size = len( title )
row_size = line_count + 1 #Extra line for title

data = [ [ None for x in range( col_size ) ] for y in range( row_size ) ]

for i in range ( col_size ):
    data[ 0 ][ i ] = title[ i ]

for j in range ( 1, row_size ):
    data[ j ][ 0 ] = Time_buf[ j - 1 ]
    data[ j ][ 1 ] = float( Mag_norm_buf[ j - 1 ] )
    data[ j ][ 2 ] = float( Mag_orth_buf[ j - 1 ] )
    data[ j ][ 3 ] = float( Board_temp_buf[ j - 1 ] )
    data[ j ][ 4 ] = float( Ambient_temp_buf[ j - 1 ] )
    data[ j ][ 5 ] = float( Humidity_buf[ j - 1 ] )
    #print ( str( data ) )

for row in data:
    ws.append(row)

chart = ScatterChart()
chart.title = "Magnitude vs Time"
chart.style = 13
chart.y_axis.title = 'Magnitude'
chart.x_axis.title = 'Time delta(seconds)'

xvalues = Reference(ws, min_col=1, min_row=2, max_row=row_size)
for i in range(2, 4):
    values = Reference(ws, min_col=i, min_row=1, max_row=row_size)
    series = Series(values, xvalues, title_from_data=True)
    chart.series.append(series)

ws.add_chart(chart, "H1")

chart2 = ScatterChart()
chart2.title = "Temperature vs Time"
chart2.style = 13 
chart2.y_axis.title = 'Degree(C)'
chart2.x_axis.title = 'Time delta(seconds)'

xvalues = Reference(ws, min_col=1, min_row=2, max_row=row_size)
values = Reference(ws, min_col=4, min_row=1, max_row=row_size)
series = Series(values, xvalues, title_from_data=True)
chart2.series.append(series)

ws.add_chart(chart2, "H20")

wb.save("Mag_Time_Temp_humidity_chart.xlsx")
